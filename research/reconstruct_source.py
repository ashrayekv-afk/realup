"""Reconstruct private eye records from the exact source workbook.

This is a new explicit parser; the prior audit's two imported parser modules were
not included in its downloadable archive. It never writes direct identifiers.
Selection rules are applied before any model is fit or evaluated.
"""
from __future__ import annotations
import argparse, hashlib, json, re
from collections import Counter
from datetime import datetime, date
from pathlib import Path
import numpy as np
import pandas as pd
from openpyxl import load_workbook

FEATURES = ["A","B","C","BAD_D","ARTmax","Kmax"]
RANGES = {"A":(2.5,12),"B":(1.5,10),"C":(150,900),"BAD_D":(-10,50),"ARTmax":(1,2000),"Kmax":(25,120)}
EXPECTED_SHA = "60bd4895f669fa55bd4e7e37266c66a52497f250847869185dfa1424bfa882be"

def numeric(v):
    if isinstance(v,(float,int)) and not isinstance(v,bool):
        return float(v) if np.isfinite(v) else np.nan
    if v is None:return np.nan
    s=str(v).strip().replace("\xa0", " ").replace(",","")
    # Legacy sheets explicitly prefix positive radius values with the label R-.
    s=re.sub(r"^R-\s*", "", s, flags=re.I)
    # A complete numeric token only; dates, mixed readings and annotations remain missing.
    if re.fullmatch(r"[+-]?(?:\d+(?:\.\d*)?|\.\d+)",s):
        return float(s)
    return np.nan

def parse_date(v, eye=None):
    if isinstance(v,(datetime,date)):
        return pd.Timestamp(v).normalize()
    if v is None:return pd.NaT
    s=str(v).strip()
    pieces = re.split(r";|\n|,",s)
    if eye and len(pieces)>1:
        explicit=[p for p in pieces if re.search(r"\b"+eye+r"\b",p,re.I)]
        if explicit: s=explicit[0]
    dates=re.findall(r"\b(\d{1,2})/(\d{1,2})/(\d{2,4})\b",s)
    if len(dates)!=1:return pd.NaT
    m,d,y=map(int,dates[0]);y=y+2000 if y<50 else y+1900 if y<100 else y
    try:return pd.Timestamp(year=y,month=m,day=d)
    except ValueError:return pd.NaT

def clean(v,feature,counters):
    n=numeric(v)
    if not np.isfinite(n):return np.nan
    lo,hi=RANGES[feature]
    if lo<=n<=hi:return n
    # No automatic repair: ambiguous units/transcription should not be guessed.
    counters["out_of_range_"+feature]+=1
    return np.nan

def run(source:Path,out:Path):
    sha=hashlib.sha256(source.read_bytes()).hexdigest()
    if sha!=EXPECTED_SHA:raise ValueError("Unrecognized source bytes; review workbook mapping before running")
    workbook=load_workbook(source,data_only=True,read_only=True)
    records=[];audit=Counter();patient_keys={}
    for sheet in workbook:
        if sheet.title not in {"Controls","CXL"}:continue
        for cells in sheet.iter_rows(min_row=2):
            values=[c.value for c in cells]
            def get(col):return values[col-1] if len(values)>=col else None
            if get(1) is None and get(2) is None:continue
            # Original linkage remains in memory. A pseudonymous index is private only.
            rawkey=str(get(2) or get(5) or get(1))
            if rawkey not in patient_keys:patient_keys[rawkey]=len(patient_keys)+1
            group=patient_keys[rawkey]
            row=cells[0].row
            note=str(get(59) or "").lower()
            legacy = sheet.title=="Controls" and get(17) is None and isinstance(get(18),str) and re.match(r"^R-\s*\d",get(18).strip(),re.I) is not None
            if legacy:audit["legacy_rows"]+=1
            mapping= {"OD":(18,35),"OS":(26,42)} if legacy else {"OD":(19,40),"OS":(29,50)}
            first_date=parse_date(get(17))
            dob=parse_date(get(4))
            reported_age=numeric(get(9))
            age=float((first_date-dob).days/365.2425) if pd.notna(first_date) and pd.notna(dob) else reported_age
            age_source="DOB_at_first_scan" if pd.notna(first_date) and pd.notna(dob) else "source_age_fallback"
            for eye,(first_start,last_start) in mapping.items():
                base={f:clean(get(first_start+i),f,audit) for i,f in enumerate(FEATURES)}
                last={f:clean(get(last_start+i),f,audit) for i,f in enumerate(FEATURES)}
                if not any(np.isfinite(v) for v in [*base.values(),*last.values()]):continue
                audit["numeric_candidate_eyes"]+=1
                excluded=[]
                if any(t in note for t in ["post-lasik","post lasik","no scans of sufficient quality","no usable scans","not able to include either eye","bad scan","poor quality","poor scan","poor qs"]):excluded.append("explicit_quality_or_diagnosis_note")
                if "blink" in note:excluded.append("recorded_blink_quality_note")
                if sheet.title=="CXL":
                    other="OS" if eye=="OD" else "OD"
                    text=" ".join(str(get(first_start+i) or "") for i in range(6)).lower()+" "+note
                    if f"only cxl {other.lower()}" in text or f"no cxl {eye.lower()}" in text or f"{eye.lower()} did not have cxl" in text:excluded.append("eye_specific_endpoint_not_documented")
                    if ("does not have keratoconus in the right eye" in note and eye=="OD") or ("does not have keratoconus in the left eye" in note and eye=="OS"):excluded.append("eye_specific_no_keratoconus")
                if not np.isfinite(age) or age<18:excluded.append("baseline_age_missing_or_under18")
                final_date=parse_date(get(38),eye)
                eye_date=parse_date(get(39 if eye=="OD" else 49),eye)
                if pd.notna(eye_date):final_date=eye_date
                years=float((final_date-first_date).days/365.2425) if pd.notna(first_date) and pd.notna(final_date) else np.nan
                # Release cohorts require a dated scan, adult age and
                # all six values from that same source scan. This stricter release
                # is not advertised as reproducing the older imputed 402-eye fit.
                # Ambiguous source-age fallback is not used in the active release.
                # The website expects age at first scan, so DOB and scan date must
                # support that same input definition in the development cohort.
                base_ok=not excluded and pd.notna(first_date) and age_source=="DOB_at_first_scan" and all(np.isfinite(v) for v in base.values())
                paired_excluded=[]
                if sheet.title=="CXL" and row==70 and eye=="OD":paired_excluded.append("documented_post_CXL_later_scan")
                if sheet.title=="CXL":
                    if "only one" in note and "scan" in note and "before" in note and "cxl" in note and not any(x in note for x in ["right eye","left eye"," od"," os","second scan before"]):paired_excluded.append("only_one_pretreatment_scan_note")
                pair_ok=base_ok and not paired_excluded and np.isfinite(years) and years>0 and all(np.isfinite(v) for v in [*base.values(),*last.values()])
                for reason in set(excluded):audit[reason]+=1
                for reason in set(paired_excluded):audit[reason]+=1
                record={"patient_id":group,"eye":eye,"outcome":int(sheet.title=="CXL"),"age_baseline":age,"interval_years":years,"age_source":age_source,"baseline_date_available":pd.notna(first_date),"baseline_eligible":base_ok,"paired_eligible":pair_ok,"source_sheet":sheet.title,"source_row":row,"legacy_layout":legacy,"exclusion_reasons":";".join(excluded+paired_excluded)}
                record.update({"baseline_"+k:v for k,v in base.items()});record.update({"latest_"+k:v for k,v in last.items()})
                records.append(record)
    frame=pd.DataFrame(records)
    baseline=frame[frame.baseline_eligible].copy()
    paired=frame[frame.paired_eligible].copy()
    def counts(d):return {"eyes":len(d),"patients":d.patient_id.nunique(),"events":int(d.outcome.sum()),"controls":int((1-d.outcome).sum()),"age_derived":int(d.age_source.eq("DOB_at_first_scan").sum()),"age_fallback":int(d.age_source.eq("source_age_fallback").sum())}
    summary={"source_sha256":sha,"parser":"New conservative explicit-layout parser; distinct from missing historical parser dependencies","candidate":counts(frame),"baseline":counts(baseline),"paired":counts(paired),"audit_counts":dict(audit),"policy":"Out-of-range values set missing; no guessed decimal repairs; complete finite pairs only; first-only baseline; no later-scan substitution; known post-CXL later scan excluded; DOB-derived baseline age required; ambiguous source-age fallback excluded."}
    out.mkdir(exist_ok=True,parents=True)
    frame.to_csv(out/"private_source_reconstruction.csv",index=False)
    baseline.to_csv(out/"private_baseline.csv",index=False)
    paired.to_csv(out/"private_paired.csv",index=False)
    (out/"reconstruction_aggregate_audit.json").write_text(json.dumps(summary,indent=2)+"\n")
    print(json.dumps(summary,indent=2))

if __name__=="__main__":
    p=argparse.ArgumentParser();p.add_argument("--source",type=Path,required=True);p.add_argument("--out",type=Path,required=True);a=p.parse_args();run(a.source,a.out)
